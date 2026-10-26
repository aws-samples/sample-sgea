#!/usr/bin/env python3
"""
Build per-baseline service-configuration catalogs from AWS sources.

Joins AWS-published artifacts:
  * sources/conformance-packs/*.yaml   — AWS Config conformance pack templates
    (awslabs/aws-config-rules, pinned commit in sources/.conformance-packs-commit):
    which Config rules apply and their parameter values.
  * sources/control-mappings/*.html    — the "Operational Best Practices for …"
    docs pages: which framework control IDs each Config rule maps to.
  * sources/artifact/extracted/        — AWS Artifact reports (NDA; ids/versions in
    sources/artifact/_downloaded.json):
      - FedRAMP SSP Appendix J CIS/CRM workbook: per-control AWS vs customer
        responsibility for FedRAMP Moderate/High and DoD IL4/IL5.
      - NIST 800-171 / CMMC 2.0 L2 CRM (GovCloud): per-practice responsibility (dod-il4).
      - AWS Customer Compliance Guides: per-service, per-security-topic customer
        guidance with control mappings (covers Amazon Bedrock).
  * sources/inferred/service-controls.yaml — SGEA-INFERRED settings for services
    AWS publishes nothing for (AgentCore, Bedrock config). Every resource and
    attribute is validated against sources/provider-schema/ (build fails on an
    unknown field); control IDs come from AWS's topic mapping and the analog
    Config rule's mapping in each baseline's pack. Marked provenance=inferred.

Writes catalog/<baseline>.json: service -> required settings (Config rules)
-> parameters -> control IDs -> source. Re-run after refreshing sources/.

Usage: python -I tools/build_catalog.py [--root /path/to/toolkit]
"""
import argparse, datetime, html, json, os, re, sys

import openpyxl

# AWS Artifact report that carries the FedRAMP SSP Appendix J CRM workbook (looked up by name in
# sources/artifact/_downloaded.json; ids and versions stay out of the source).
FEDRAMP_CRM_REPORT = "FedRAMP Customer Package"
import yaml

DOCS = "https://docs.aws.amazon.com/config/latest/developerguide/operational-best-practices-for-{}.html"
PACK_URL = ("https://github.com/awslabs/aws-config-rules/blob/{sha}/aws-config-conformance-packs/"
            "Operational-Best-Practices-for-{}.yaml")

# Pack -> docs control-mapping slug. WHICH packs a baseline uses is a control
# decision and lives in baselines/<id>.json evidence.conformance_packs.
PACK_DOCS = {
    "FedRAMP-Low": "fedramp-low", "FedRAMP-Moderate": "fedramp-moderate",
    "FedRAMP-HighPart1": "fedramp-high-part-1", "FedRAMP-HighPart2": "fedramp-high-part-2",
    "HIPAA-Security": "hipaa_security", "FDA-21CFR-Part-11": "FDA-21CFR-Part-11",
    "NIST-800-53-rev-5": "nist-800-53_rev_5", "CMMC-2.0-Level-2": "cmmc_2.0_level_2",
}
# Baselines with no official AWS pack, where a NIST-derived pack stands in.
PROXY_NOTES = {
    "dod-il4": "No AWS conformance pack exists for DoD IL4 (DISA CC SRG). NIST 800-53 Rev5 + "
               "CMMC 2.0 L2 stand in; IL4-specific SRG controls (e.g. CNSSI 1253 overlays) are NOT covered.",
    "dod-il2": "No AWS conformance pack exists for the DoD SRG. FedRAMP Moderate (composed) + CMMC 2.0 L2 stand in.",
    "dod-il5": "No AWS conformance pack exists for the DoD SRG. FedRAMP High + CMMC 2.0 L2 stand in; IL5-specific "
               "SRG controls (CNSSI 1253 overlays, SCCA) are NOT covered by Config rules.",
    "hitrust": "No AWS conformance pack exists for HITRUST CSF. NIST 800-53 Rev5 stands in; "
               "HITRUST requirement IDs are NOT mapped.",
}


def baselines(root):
    """baseline id -> {packs: [(pack, docs slug)], proxy} from baselines/*.json."""
    out = {}
    for f in sorted(os.listdir(os.path.join(root, "baselines"))):
        if f.endswith(".json"):
            bid = f[:-5]
            frag = json.load(open(os.path.join(root, "baselines", f)))
            packs = frag.get("evidence", {}).get("conformance_packs", [])
            out[bid] = {"packs": [(p, PACK_DOCS[p]) for p in packs], "proxy": PROXY_NOTES.get(bid)}
    return out


ARTIFACT = os.path.join("artifact", "extracted")
CRM_XLSX = os.path.join(ARTIFACT, "FedRAMP_Customer_Package",
                        "AWS_SSP-Appendix-J_CIS-and-CRM-Workbook_v26.01.U-FINAL.xlsx")
CMMC_XLSX = os.path.join(ARTIFACT, "AWS_NIST_SP_800-171___CMMC_Customer_Package",
                         "AWS NIST SP 800-171 CMMC 2.0 Level 2 Customer Responsibility Matrix v2.5.xlsx")
GUIDE_XLSX = os.path.join(ARTIFACT, "Customer_Compliance_Guide", "AWS Customer Compliance Guides.xlsx")

# Which CRM rows belong to each baseline. The CRM designates rows as
# "FedRAMP High & Mod" / "FedRAMP High" and adds DoD rows on top of FedRAMP.
CRM_SCOPE = {
    "fedramp-low":      {"fedramp": ("Mod",), "nist_profile": "LOW",
                         "note": "The AWS CRM has no Low designation: Moderate rows are trimmed to the NIST "
                         "SP 800-53B LOW baseline (usnistgov/oscal-content, pinned in sources/.nist-oscal-commit). "
                         "FedRAMP Low adds a few parameters/controls on top of NIST Low; confirm at ISSO review."},
    "fedramp-moderate": {"fedramp": ("Mod",)},
    "fedramp-high":     {"fedramp": ("High",)},
    "dod-il2":          {"fedramp": ("Mod",), "note": "DoD IL2 adds no CRM rows beyond FedRAMP Moderate."},
    "dod-il4":          {"fedramp": ("Mod",), "dod": ("DoD IL4",)},
    "dod-il5":          {"fedramp": ("High",), "dod": ("IL5",)},
}
# Compliance Guide framework column(s) whose control mappings apply per baseline.
GUIDE_FRAMEWORKS = {
    "fedramp-low": ["NIST 800-53 Rev5"], "fedramp-moderate": ["NIST 800-53 Rev5"],
    "fedramp-high": ["NIST 800-53 Rev5"], "hipaa": ["HIPAA", "NIST 800-53 Rev5"],
    "fda-21-cfr-part-11": ["NIST 800-53 Rev5"], "hitrust": ["NIST 800-53 Rev5"],
    "dod-il4": ["NIST 800-53 Rev5", "CMMC"],
    "dod-il2": ["NIST 800-53 Rev5", "CMMC"], "dod-il5": ["NIST 800-53 Rev5", "CMMC"],
}
# SGEA services.allow ids (plus enclave foundation services) -> Compliance Guide service.
SGEA_GUIDE_SERVICES = {
    "bedrock": "Amazon Bedrock", "bedrock-agentcore": None,
    "aurora-postgresql": "Amazon Aurora", "kms": "AWS Key Management Service (KMS)",
    "cloudwatch": "Amazon CloudWatch + EventBridge", "eventbridge": "Amazon CloudWatch + EventBridge",
    "cloudtrail": "AWS CloudTrail", "s3": "Amazon S3 and Glacier", "lambda": "AWS Lambda",
    "dynamodb": "Amazon DynamoDB", "sns": "Amazon Simple Notification Service (SNS)",
    "firehose": "Amazon Data Firehose", "cognito": "Amazon Cognito",
    "elasticache": "Amazon ElastiCache for Redis",
    "iam": "AWS Identity and Access Management (IAM)", "organizations": "AWS Organizations",
    "vpc": "Amazon Virtual Private Cloud (VPC)", "config": "AWS Config",
    "guardduty": "Amazon GuardDuty", "secretsmanager": "AWS Secrets Manager",
    "dod-scca": None,
}

ROW = re.compile(r"<tr>(.*?)</tr>", re.S)
CELL = re.compile(r"<td[^>]*>(.*?)</td>", re.S)
LINK = re.compile(r"<a [^>]*>([^<]+)</a>")
TAG = re.compile(r"<[^>]+>")


def _text(s):
    return re.sub(r"\s+", " ", html.unescape(TAG.sub(" ", s))).strip()


def parse_mappings(path):
    """rule name -> [{id, description}] from a docs mapping table."""
    out = {}
    for row in ROW.findall(open(path, encoding="utf-8").read()):
        cells = CELL.findall(row)
        if len(cells) < 3:
            continue
        cid, desc = _text(cells[0]), _text(cells[1])
        for rule in LINK.findall(cells[2]) or [_text(cells[2])]:
            rule = rule.strip()
            entry = {"id": cid, "description": desc}
            if entry not in out.setdefault(rule, []):
                out[rule].append(entry)
    return out


def _param(value, params):
    """Resolve Fn::If/Ref parameter indirection to the pack's default value."""
    if isinstance(value, dict):
        if "Fn::If" in value:
            return _param(value["Fn::If"][1], params)
        if "Ref" in value:
            return params.get(value["Ref"], {}).get("Default")
    return value


def parse_pack(path):
    """rule name -> rule details from a conformance pack template."""
    doc = yaml.safe_load(open(path, encoding="utf-8"))
    params = doc.get("Parameters", {}) or {}
    rules = {}
    for res in (doc.get("Resources") or {}).values():
        if res.get("Type") != "AWS::Config::ConfigRule":
            continue
        p = res["Properties"]
        src = p.get("Source", {})
        types = (p.get("Scope") or {}).get("ComplianceResourceTypes", [])
        inputs = {k: _param(v, params) for k, v in (p.get("InputParameters") or {}).items()}
        rules[p["ConfigRuleName"]] = {
            "config_rule": p["ConfigRuleName"],
            "owner": src.get("Owner"),
            "source_identifier": src.get("SourceIdentifier"),
            "custom_policy": (src.get("CustomPolicyDetails") or {}).get("PolicyText"),
            "resource_types": types,
            "parameters": {k: v for k, v in inputs.items() if v is not None},
        }
    return rules


# Normalize Config-rule prefixes / CFN service names onto one bucket per service.
SERVICE_ALIASES = {
    "access": "iam", "mfa": "iam", "root": "iam", "annual": "account",
    "cloud": "cloudtrail", "multi": "cloudtrail", "alb": "elb", "elbv2": "elb",
    "elasticloadbalancing": "elb", "elasticloadbalancingv2": "elb",
    "elasticsearch": "opensearch", "vpc": "ec2", "restricted": "ec2", "instances": "ec2",
    "ebs": "ec2", "db": "rds", "aurora": "rds",
}


CONTROL = re.compile(r"^([A-Z]{2})-0*(\d+)(?:\s*\(0*(\d+)\))?")


def control_key(cid):
    """'AC-02(03).a' / 'AC-2(3)' -> 'AC-2(3)'; non-NIST ids (HIPAA, CMMC) -> None."""
    m = CONTROL.match((cid or "").strip().upper())
    if not m:
        return None
    return m.group(1) + "-" + m.group(2) + ("(" + m.group(3) + ")" if m.group(3) else "")


def _rows(path, sheet, skip):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return [r for r in list(wb[sheet].iter_rows(values_only=True))[skip:] if r and r[0]]


def _s(v):
    return str(v).strip() if v is not None else None


def _lines(v):
    return [x.strip() for x in re.split(r"[\n,]", _s(v) or "") if x.strip()]


CIS_COLS = ["implemented", "partially_implemented", "planned", "alternative_implementation", "n/a",
            "service_provider_corporate", "service_provider_system_specific", "service_provider_hybrid",
            "configured_by_customer", "provided_by_customer", "shared", "inherited_from_pre_existing_authorization"]


def load_crm(srcdir):
    """FedRAMP/DoD CRM rows, with CIS implementation status + origination merged in."""
    cis = {}
    for r in _rows(os.path.join(srcdir, CRM_XLSX), "Combined CIS - FedRAMP and DoD", 3):
        cis[_s(r[0])] = [c for c, v in zip(CIS_COLS, r[3:15]) if _s(v) == "X"]
    out = []
    for r in _rows(os.path.join(srcdir, CRM_XLSX), "Combined CRM - FedRAMP and DoD", 3):
        cid = _s(r[0])
        flags = cis.get(cid, [])
        out.append({
            "control_id": cid, "control_key": control_key(cid),
            "fedramp_baseline": _s(r[1]), "dod_baseline": _s(r[2]),
            "can_be_inherited_from_aws": _s(r[3]),
            "customer_responsibility": _s(r[4]),
            "implementation_status": [f for f in flags if f in CIS_COLS[:5]],
            "control_origination": [f for f in flags if f not in CIS_COLS[:5]],
        })
    return out


NIST_OSCAL = os.path.join("nist-oscal", "NIST_SP-800-53_rev5_{}-baseline_profile.json")


def nist_baseline_keys(srcdir, level):
    """control keys ('AC-2', 'AC-2(1)') in a NIST SP 800-53B OSCAL baseline profile."""
    d = json.load(open(os.path.join(srcdir, NIST_OSCAL.format(level))))
    ids = [i for imp in d["profile"]["imports"] for inc in imp.get("include-controls", []) for i in inc.get("with-ids", [])]
    out = set()
    for i in ids:
        base, _, enh = i.upper().partition(".")
        out.add(control_key(base + ("(%s)" % enh if enh else "")))
    return out


def crm_for(baseline, crm, srcdir=None):
    scope = CRM_SCOPE.get(baseline)
    if not scope:
        return None
    keep = nist_baseline_keys(srcdir, scope["nist_profile"]) if scope.get("nist_profile") and srcdir else None
    def hit(v, needles):
        return bool(v) and any(n in v for n in needles)
    rows = [r for r in crm if hit(r["fedramp_baseline"], scope.get("fedramp", ()))
            or hit(r["dod_baseline"], scope.get("dod", ()))]
    return [r for r in rows if r["control_key"] in keep] if keep is not None else rows


def load_cmmc(srcdir):
    out = []
    for r in _rows(os.path.join(srcdir, CMMC_XLSX), "CMMC CRM GovCloud (US)", 1):
        out.append({
            "domain": _s(r[0]), "practice_ids": _lines(r[1]), "nist_800_171": _lines(r[2]),
            "nist_800_53": _lines(r[3]), "practice": _s(r[4]), "fulfilled_by": _s(r[5]),
            "can_be_inherited_from_aws": _s(r[6]), "aws_services": _lines(r[7]),
            "aws_implementation": _s(r[8]), "customer_responsibility": _s(r[9]),
        })
    return out


def _topic(t):
    return re.sub(r"\s+", " ", (t or "").strip()).lower()


def load_guide(srcdir):
    """service -> [topic rows]; 'Related Controls' is an Excel VLOOKUP, resolved here."""
    path = os.path.join(srcdir, GUIDE_XLSX)
    mrows = list(openpyxl.load_workbook(path, read_only=True)["Mapping Rows Merged_2024"].iter_rows(values_only=True))
    header = [_s(h) for h in mrows[0]]
    mapping = {}
    for r in mrows[1:]:
        if r[0]:
            mapping[_topic(r[0])] = {header[i]: [c for c in _lines(r[i]) if c != "No Control Mapping"]
                                     for i in range(1, len(header)) if header[i] and header[i] != "None"}
    guide = {"_topic_mapping": mapping}
    for r in _rows(path, "AWS Customer Compliance Guides", 1):
        guide.setdefault(_s(r[0]), []).append({
            "topic": _s(r[1]), "guidance": _s(r[3]), "supporting_services": _lines(r[4]),
            "controls": mapping.get(_topic(r[1]), {}),
        })
    return guide


# ---- Services in Scope (which services are authorized at each level) ----
SIS_FILES = {"fedramp": os.path.join("public", "fedramp-services-in-scope.html"),
             "dod": os.path.join("public", "dod-ccsrg-services-in-scope.html"),
             "hipaa": os.path.join("public", "hipaa-eligible-services-reference.html")}
# SGEA services.allow id -> service names used on the AWS Services-in-Scope pages.
SIS_NAMES = {
    "bedrock": ["Amazon Bedrock"], "bedrock-agentcore": ["Amazon Bedrock AgentCore"],
    "aurora-postgresql": ["Amazon Aurora PostgreSQL", "Amazon Aurora"],
    "kms": ["AWS Key Management Service"], "cloudwatch": ["Amazon CloudWatch"],
    "cloudtrail": ["AWS CloudTrail"], "s3": ["Amazon Simple Storage Service", "Amazon S3"],
    "lambda": ["AWS Lambda"], "dynamodb": ["Amazon DynamoDB"],
    "sns": ["Amazon Simple Notification Service", "Amazon SNS"], "eventbridge": ["Amazon EventBridge"],
    "firehose": ["Amazon Data Firehose", "Amazon Kinesis Data Firehose"], "cognito": ["Amazon Cognito"],
    "elasticache": ["Amazon ElastiCache"],
}


def _sis_key(name):
    return re.sub(r"[^a-z0-9]", "", re.sub(r"\(.*?\)|\*", "", name.lower()))


def load_services_in_scope(srcdir):
    """{program column -> {normalized service name -> status}} from the public pages."""
    out = {}
    for prog in ("fedramp", "dod"):
        t = open(os.path.join(srcdir, SIS_FILES[prog]), encoding="utf-8").read()
        tab = re.findall(r"<table.*?</table>", t, re.S)[0]
        rows = [[_text(c) for c in re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, re.S)]
                for r in re.findall(r"<tr.*?</tr>", tab, re.S)]
        header = [re.sub(r"\\n", " ", h) for h in rows[0]]
        for r in rows[1:]:
            for col, val in zip(header[2:], r[2:]):
                out.setdefault(col, {})[_sis_key(r[0])] = val or ""
    t = open(os.path.join(srcdir, SIS_FILES["hipaa"]), encoding="utf-8").read()
    body = _text(t[t.find("HIPAA Eligible Services"):])
    out["HIPAA Eligible"] = {_sis_key(n): "\u2713" for names in SIS_NAMES.values() for n in names
                             if re.search(r"(^|\s)" + re.escape(n) + r"(\s|$)", body)}
    return out


def services_status(sis):
    """SGEA service id -> {program column -> status} (empty string = not listed / not authorized)."""
    res = {}
    for sid, names in SIS_NAMES.items():
        res[sid] = {}
        for col, table in sis.items():
            vals = [table.get(_sis_key(n)) for n in names if table.get(_sis_key(n)) is not None]
            res[sid][col] = vals[0] if vals else ""
    return res


# ---- Region availability (AWS Regional Services List) + GovCloud differences ----
REGIONAL_TABLE = os.path.join("public", "guides", "regional-services-table.json")
REGION_NAMES = {
    "bedrock": "Amazon Bedrock", "bedrock-agentcore": "Amazon Bedrock AgentCore",
    "aurora-postgresql": "Amazon Aurora PostgreSQL", "kms": "AWS Key Management Service (KMS)",
    "cloudwatch": "Amazon CloudWatch", "cloudtrail": "AWS CloudTrail",
    "s3": "Amazon Simple Storage Service (S3)", "lambda": "AWS Lambda", "dynamodb": "Amazon DynamoDB",
    "sns": "Amazon Simple Notification Service (SNS)", "eventbridge": "Amazon EventBridge",
    "firehose": "Amazon Data Firehose", "cognito": "Amazon Cognito", "elasticache": "Amazon ElastiCache",
    "config": "AWS Config",
}
GOVCLOUD_PAGES = {"bedrock": "govcloud-govcloud-bedrock.html", "bedrock-agentcore": "govcloud-govcloud-bedrock-agentcore.html",
                  "aurora-postgresql": "govcloud-govcloud-aurora.html"}
MODELS_PAGE = os.path.join("public", "guides", "bedrock-models-fedramp-dod.html")


def service_regions(srcdir):
    d = json.load(open(os.path.join(srcdir, REGIONAL_TABLE)))
    by = {}
    for p in d["prices"]:
        a = p["attributes"]
        by.setdefault(a.get("aws:serviceName"), set()).add(a.get("aws:region"))
    return {"source_version": d.get("metadata", {}).get("source:version"),
            "services": {sid: sorted(by.get(name, [])) for sid, name in REGION_NAMES.items()}}


def govcloud_differences(srcdir):
    """Per service: GovCloud regions, CloudFormation types AWS lists as not available, other differences."""
    out = {}
    for sid, page in GOVCLOUD_PAGES.items():
        t = _text(open(os.path.join(srcdir, "public", "guides", page), encoding="utf-8").read())
        avail = t.rsplit("Region availability", 1)[-1].split("differs", 1)[0]  # last = section body, not on-page TOC
        regions = sorted({"us-gov-" + m.lower() + "-1" for m in re.findall(r"US-(West|East)", avail)})
        diffs = t.split("differences apply", 1)[-1].split("Documentation", 1)[0]
        unavailable = sorted(set(re.findall(r"AWS::[A-Za-z0-9]+::[A-Za-z0-9]+", diffs)))
        notes = sorted({n.strip() for n in re.findall(r"([A-Z][A-Za-z ]+? (?:is|are) not availa\w+)", diffs)
                        if "CloudFormation" not in n})
        out[sid] = {"regions": regions, "unavailable_cfn_types": unavailable, "notes": notes}
    return out


MODEL_IDS = os.path.join("public", "guides", "bedrock-list-foundation-models.json")


# Tokens an API model name may carry beyond the authorization-table name without being a
# different model: vendor/brand words and build qualifiers. Any other extra token (a version
# number, "safeguard", "plus", ...) means a DIFFERENT model and blocks the match.
MODEL_NAME_NOISE = {"nvidia", "openai", "amazon", "anthropic", "meta", "xai", "ai21", "labs", "titan",
                    "v1", "v2", "bf16", "vl", "a12b"}
CURATED_MODEL_IDS = os.path.join("inferred", "bedrock-model-ids.json")


def _model_tokens(name):
    # word tokens keep trailing digits ("v2", "g1", "bf16", "a12b"); version tokens keep a unit ("70b", "5.5")
    toks = re.findall(r"[a-z]+\d*[a-z]*\+?|\d+(?:\.\d+)?[a-z]*", name.lower().replace("-", " "))
    return {re.sub(r"^(\d+)\.0$", r"\1", t) for t in toks} - {"and", "the"}   # "5.0" == "5"


def model_ids_for(name, api):
    """Exact model-name match (modulo MODEL_NAME_NOISE) against ListFoundationModels names.
    'GPT 5.6 (Terra and Luna)' -> one exact match per variant. Exact IDs only, never prefixes."""
    variants = [name]
    m = re.match(r"(.*?)\s*\((.+)\)", name)
    if m:
        variants = [m.group(1) + " " + v for v in re.split(r"\s+and\s+|,\s*", m.group(2))]
    ids = set()
    for v in variants:
        want = _model_tokens(v) - MODEL_NAME_NOISE
        for x in api:
            if _model_tokens(x["modelName"]) - MODEL_NAME_NOISE == want:
                ids.add(x["modelId"])
    return sorted(ids)


def bedrock_models(srcdir):
    """Per-model authorization from AWS 'Amazon Bedrock models — FedRAMP and DoD CSP SRG' page."""
    t = open(os.path.join(srcdir, MODELS_PAGE), encoding="utf-8").read()
    by = {}
    for cls, body in re.findall(r'<td class="[^"]*?\s([^"]+?-\d+)" data-rg-n="TableCell">(.*?)</td>', t, re.S):
        by.setdefault(cls, []).append(_text(body))
    rows = []
    for cls, vals in by.items():
        fam = cls.rsplit("-", 1)[0].split(" ")[-1] if " " not in cls.rsplit("-", 1)[0].split(" ", 1)[-1] \
            else cls.rsplit("-", 1)[0].split(" ", 1)[1]
        vals = (vals + ["", "", "", ""])[:4]
        rows.append({"family": fam, "model": vals[0], "fedramp_class_c": vals[1] == "X",
                     "fedramp_class_d_govcloud": vals[2] == "X", "dod_il4_il5": vals[3] == "X"})
    path = os.path.join(srcdir, MODEL_IDS)
    api = [] if not os.path.exists(path) else [m for ms in json.load(open(path)).values() for m in ms]
    api = list({m["modelId"]: m for m in api}.values())
    curated = json.load(open(os.path.join(srcdir, CURATED_MODEL_IDS)))["models"]
    for r in rows:
        r["model_ids"] = model_ids_for(r["model"], api)
        r["model_id_source"] = "ListFoundationModels" if r["model_ids"] else None
        if not r["model_ids"] and (r["family"] + " " + r["model"]) in curated:
            r["model_ids"] = curated[r["family"] + " " + r["model"]]
            r["model_id_source"] = "sgea-curated (verify)"
    return rows


INFERRED = os.path.join("inferred", "service-controls.yaml")
SCHEMA_SLICE = os.path.join("provider-schema", "bedrock-agentcore-slice.json")
PLAN_CHECKS = os.path.join("inferred", "config-rule-plan-checks.json")


def _schema_has(block, path):
    """Walk a Terraform block schema (attributes, nested_type, block_types) by dotted path."""
    for part in path.split("."):
        attrs = block.get("attributes", {})
        if part in attrs:
            nt = attrs[part].get("nested_type")
            block = nt if nt else {}
        elif part in block.get("block_types", {}):
            block = block["block_types"][part]["block"]
        else:
            return False
    return True


def load_inferred(srcdir, guide):
    doc = yaml.safe_load(open(os.path.join(srcdir, INFERRED)))
    schema = json.load(open(os.path.join(srcdir, SCHEMA_SLICE)))
    topics = guide["_topic_mapping"]
    errs = []
    for c in doc["controls"]:
        res = schema["resource_schemas"].get(c["resource"])
        if res is None:
            errs.append("%s: unknown resource %s" % (c["id"], c["resource"]))
        else:
            for attr in ([c["attribute"]] if c.get("attribute") else []) + list(c.get("any_of") or []):
                if not _schema_has(res, attr):
                    errs.append("%s: %s has no attribute %s" % (c["id"], c["resource"], attr))
        if _topic(c["topic"]) not in topics:
            errs.append("%s: unknown Compliance Guide topic %r" % (c["id"], c["topic"]))
    plan_checks = json.load(open(os.path.join(srcdir, PLAN_CHECKS)))
    for rule, checks in plan_checks["rules"].items():
        for c in checks:
            for rtype in {c["resource"], c.get("subject", c["resource"])}:
                if rtype not in schema["resource_schemas"]:
                    errs.append("plan-check %s: unknown resource %s" % (rule, rtype))
            res = schema["resource_schemas"].get(c["resource"])
            if res and c.get("attribute") and not _schema_has(res, c["attribute"]):
                errs.append("plan-check %s: %s has no attribute %s" % (rule, c["resource"], c["attribute"]))
    if errs:
        raise SystemExit("inferred controls failed validation:\n  " + "\n  ".join(errs))
    return doc["controls"], schema["providers"]


def _lookup(frag, path):
    for k in path.split("."):
        if not isinstance(frag, dict) or k not in frag:
            return None
        frag = frag[k]
    return frag


def inferred_for(baseline, inferred, guide, rules, frag):
    """Resolve inferred controls' control IDs for this baseline's frameworks."""
    fws = GUIDE_FRAMEWORKS[baseline]
    topics = guide["_topic_mapping"]
    out = {}
    for c in inferred:
        controls = {fw: list(topics[_topic(c["topic"])].get(fw, [])) for fw in fws}
        for fw, ids in (c.get("extra_controls") or {}).items():
            if fw in controls:
                controls[fw] += [i for i in ids if i not in controls[fw]]
        analog = rules.get(c.get("analog_config_rule"))
        entry = {k: c[k] for k in ("id", "resource", "attribute", "require", "level", "rationale",
                                   "topic", "analog_config_rule", "applies_when", "value_from", "scope", "where", "cfn", "any_of") if c.get(k)}
        if c.get("value_from"):
            entry["resolved_value"] = _lookup(frag, c["value_from"])
        entry.update({
            "provenance": "inferred",
            "controls_from_topic": controls,
            "controls_from_analog_rule": [{"id": x["id"], "framework": x["framework"]}
                                          for x in (analog or {}).get("controls", [])],
            "analog_rule_in_baseline_pack": bool(analog) if c.get("analog_config_rule") else None,
        })
        out.setdefault(c["service"], []).append(entry)
    return out


def guidance_for(baseline, guide):
    fws = GUIDE_FRAMEWORKS[baseline]
    out = {}
    for sgea, name in SGEA_GUIDE_SERVICES.items():
        if not name or name not in guide:
            out[sgea] = {"guide_service": None, "topics": [],
                         "note": "No AWS Customer Compliance Guide exists for this service."}
            continue
        out[sgea] = {"guide_service": name, "topics": [
            dict(t, controls={fw: t["controls"].get(fw, []) for fw in fws}) for t in guide[name]]}
    return out


def _service(rule):
    if rule["resource_types"]:
        svc = rule["resource_types"][0].split("::")[1].lower()
    else:
        svc = rule["config_rule"].split("-")[0]
    return SERVICE_ALIASES.get(svc, svc)


# Verbatim text from NDA-covered AWS Artifact reports. Kept out of catalog/ (which may be
# shared): only IDs, classifications and mappings are written there. `--include-artifact-text`
# writes the full documents to catalog/_nda/ for internal use.
ARTIFACT_TEXT_KEYS = {"customer_responsibility", "guidance", "aws_implementation"}
GEN_START = "<!-- SGEA:GENERATED START (tools/build_catalog.py — do not edit by hand) -->"
GEN_END = "<!-- SGEA:GENERATED END -->"


def _redact(obj):
    if isinstance(obj, dict):
        return {k: ("[AWS Artifact text omitted — see artifact_sources]" if k in ARTIFACT_TEXT_KEYS and v else _redact(v))
                for k, v in obj.items()}
    if isinstance(obj, list):
        return [_redact(x) for x in obj]
    return obj


def _write(path, doc):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w") as f:
        json.dump(doc, f, indent=2)
        f.write("\n")


def _family(key):
    return (key or "").split("-")[0]


def responsibility_section(cat):
    """Markdown facts for verticals/<baseline>/responsibility-matrix.md (no AWS text)."""
    lines = [GEN_START, "", "## Generated responsibility summary", ""]
    src = [a for a in cat["artifact_sources"] if a["name"] == FEDRAMP_CRM_REPORT]
    # Control keys the enclave IaC addresses: mapped Config rules + inferred controls.
    iac = set()
    for rs in cat["services"].values():
        for r in rs:
            iac |= {control_key(c["id"]) for c in r["controls"] if control_key(c["id"])}
    for svc in cat["service_guidance"].values():
        for ic in svc.get("inferred_controls", []):
            for ids in ic["controls_from_topic"].values():
                iac |= {control_key(i) for i in ids if control_key(i)}
    if cat["crm"]:
        rows = cat["crm"]
        lines += ["Source: AWS Artifact *%s* (report `%s`, version %s) — SSP Appendix J CRM. Classifications "
                  "only; AWS's responsibility text is in the report, not reproduced here."
                  % (src[0]["name"], src[0]["id"], src[0]["version"]) if src else "Source: AWS CRM.", ""]
        if cat.get("crm_note"):
            lines += ["> " + cat["crm_note"], ""]
        cnt = {}
        for r in rows:
            cnt[r["can_be_inherited_from_aws"] or "unspecified"] = cnt.get(r["can_be_inherited_from_aws"] or "unspecified", 0) + 1
        lines += ["| Can be inherited from AWS | Control parts |", "|---|---|"]
        lines += ["| %s | %d |" % (k, v) for k, v in sorted(cnt.items())] + [""]
        fam = {}
        for r in rows:
            if r["can_be_inherited_from_aws"] in ("Partial", "No"):
                f = fam.setdefault(_family(r["control_key"]), {"Partial": 0, "No": 0, "iac": 0})
                f[r["can_be_inherited_from_aws"]] += 1
                f["iac"] += 1 if r["control_key"] in iac else 0
        lines += ["Customer-side control parts by family, and how many the enclave IaC addresses "
                  "(a mapped Config rule or SGEA control checked by `sgea check`); the rest are "
                  "policy/process items for the platform team, tenant or ISSO.", "",
                  "| Family | Partial | No (customer) | Addressed by enclave IaC |", "|---|---|---|---|"]
        lines += ["| %s | %d | %d | %d |" % (k, v["Partial"], v["No"], v["iac"]) for k, v in sorted(fam.items())]
    else:
        lines += ["Overlay baseline — no framework-specific AWS CRM; responsibility for the underlying "
                  "NIST controls comes from the FedRAMP baseline the profile composes it with.", ""]
        secs = {}
        for rs in cat["services"].values():
            for r in rs:
                for c in r["controls"]:
                    if not control_key(c["id"]):
                        secs.setdefault(c["id"], set()).add(r["config_rule"])
        if secs:
            lines += ["Framework requirements with automated evidence (AWS Config rules in the composed "
                      "conformance pack, deployed by SGEA):", "",
                      "| Requirement | Config rules |", "|---|---|"]
            lines += ["| %s | %d |" % (k, len(v)) for k, v in sorted(secs.items())]
    lines += ["", GEN_END]
    return "\n".join(lines) + "\n"


def update_vertical(root, baseline, section):
    path = os.path.join(root, "verticals", baseline, "responsibility-matrix.md")
    if not os.path.exists(path):
        return
    text = open(path).read()
    if GEN_START in text:
        text = text[:text.index(GEN_START)] + section + text[text.index(GEN_END) + len(GEN_END) + 1:]
    else:
        text = text.rstrip("\n") + "\n\n" + section
    open(path, "w").write(text)


DISPOSITIONS = ("implement", "tenant", "not-in-enclave", "customer-prerequisite", "process", "runtime-only")


def load_coverage(srcdir):
    """(checked rule names, dispositions). A rule is 'checked' when sgea check evaluates it on both the
    Terraform plan and the CDK template (config-rule-plan-checks.json rules + cfn_rules)."""
    import yaml
    pc = json.load(open(os.path.join(srcdir, "inferred", "config-rule-plan-checks.json")))
    checked = set(pc["rules"]) & set(pc["cfn_rules"])
    disp = (yaml.safe_load(open(os.path.join(srcdir, "inferred", "rule-dispositions.yaml"))) or {}).get("dispositions") or {}
    errs = ["%s: disposition %r not in %s" % (n, d.get("disposition"), DISPOSITIONS)
            for n, d in disp.items() if d.get("disposition") not in DISPOSITIONS or not d.get("reason")]
    errs += ["%s: has a plan/CFN check AND a disposition (remove the disposition)" % n for n in sorted(set(disp) & checked)]
    errs += ["%s: plan check without a CFN check (or vice versa)" % n
             for n in sorted(set(pc["rules"]) ^ set(pc["cfn_rules"]))]
    if errs:
        raise SystemExit("rule dispositions failed validation:\n  " + "\n  ".join(errs))
    return checked, disp


def coverage_for(rules, checked, disp):
    cov, missing = {}, []
    for name in sorted(rules):
        if name in checked:
            cov[name] = {"status": "checked"}
        elif name in disp:
            cov[name] = {"status": disp[name]["disposition"], "reason": disp[name]["reason"]}
        else:
            missing.append(name)
    return cov, missing


def coverage_report(root, per_baseline, disp):
    """docs/CONTROL-COVERAGE.md: every Config rule of every composed pack, per baseline."""
    order = ("checked",) + DISPOSITIONS
    lines = ["# SGEA control coverage (generated by tools/build_catalog.py — do not edit)", "",
             "Every AWS Config rule in every conformance pack each baseline composes, and how SGEA accounts "
             "for it. `checked` = evaluated by `sgea check` on the Terraform plan and the CDK template; "
             "`implement` = enclave-owned, still to be emitted (an open gap: `sgea check` reports it as GAP). "
             "Reasons: `sources/inferred/rule-dispositions.yaml`. SGEA's Bedrock/AgentCore controls are listed in "
             "each catalog (`service_guidance.*.inferred_controls`) and are all checked.", "",
             "| baseline | rules | " + " | ".join(order) + " |", "|---|---:|" + "---:|" * len(order)]
    for b, cov in sorted(per_baseline.items()):
        cnt = {k: sum(1 for v in cov.values() if v["status"] == k) for k in order}
        lines.append("| %s | %d | %s |" % (b, len(cov), " | ".join(str(cnt[k]) for k in order)))
    allrules = {}
    for b, cov in per_baseline.items():
        for n, v in cov.items():
            allrules.setdefault(n, dict(v, baselines=[]))["baselines"].append(b)
    for k in order:
        names = sorted(n for n, v in allrules.items() if v["status"] == k)
        if not names:
            continue
        lines += ["", "## %s (%d)" % (k, len(names)), "", "| rule | baselines | reason |", "|---|---|---|"]
        for n in names:
            lines.append("| `%s` | %s | %s |" % (n, ", ".join(sorted(allrules[n]["baselines"])),
                                                 allrules[n].get("reason", "")))
    import yaml
    excs = (yaml.safe_load(open(os.path.join(root, "sources", "inferred", "rule-dispositions.yaml"))) or {}).get("exceptions") or []
    unknown = sorted({x["rule"] for x in excs} - set(allrules))
    if unknown:
        raise SystemExit("rule-dispositions.yaml exceptions name rules no composed pack contains: " + ", ".join(unknown))
    if excs:
        lines += ["", "## Documented exceptions (%d)" % len(excs), "",
                  "Enclave resources AWS Config reports NON_COMPLIANT by design. The sandbox evidence "
                  "(`docs/evidence/`) marks matching live findings as exceptions, with these reasons.", "",
                  "| rule | resource | reason |", "|---|---|---|"]
        lines += ["| `%s` | `%s` | %s |" % (x["rule"], x["resource"], x["reason"]) for x in excs]
    os.makedirs(os.path.join(root, "docs"), exist_ok=True)
    with open(os.path.join(root, "docs", "CONTROL-COVERAGE.md"), "w") as f:
        f.write("\n".join(lines) + "\n")


def build(root, include_artifact_text=False):
    srcdir = os.path.join(root, "sources")
    sha = open(os.path.join(srcdir, ".conformance-packs-commit")).read().strip()
    outdir = os.path.join(root, "catalog")
    os.makedirs(outdir, exist_ok=True)
    summary = []
    crm, cmmc, guide = load_crm(srcdir), load_cmmc(srcdir), load_guide(srcdir)
    inferred, schema_providers = load_inferred(srcdir, guide)
    checked, disp = load_coverage(srcdir)
    per_baseline, missing_all = {}, {}
    downloaded = json.load(open(os.path.join(srcdir, "artifact", "_downloaded.json")))
    with open(os.path.join(outdir, "_services-in-scope.json"), "w") as f:
        json.dump({"_doc": "AWS Services in Scope (sources/public/*.html): authorization status per SGEA "
                           "service id and program column. '\u2713' = authorized; any other text (e.g. "
                           "'Independent Assessment') or '' = NOT authorized at that level. Consumed by "
                           "sgea validate.",
                   "services": services_status(load_services_in_scope(srcdir))}, f, indent=2)
        f.write("\n")
    _write(os.path.join(outdir, "_regions.json"), {
        "_doc": "Region facts for sgea validate: service availability per region (AWS Regional Services "
                "List, sources/public/guides/regional-services-table.json), GovCloud service differences "
                "(GovCloud User Guide pages), and per-model Bedrock authorization (AWS 'Amazon Bedrock models "
                "- FedRAMP and DoD CSP SRG' page).",
        "service_regions": service_regions(srcdir),
        "govcloud_differences": govcloud_differences(srcdir),
        "bedrock_models": bedrock_models(srcdir)})
    crm_doc = {"source": [d for d in downloaded if d["name"] == FEDRAMP_CRM_REPORT], "controls": crm}
    _write(os.path.join(outdir, "_crm-fedramp-dod.json"), _redact(crm_doc))
    if include_artifact_text:
        _write(os.path.join(outdir, "_nda", "_crm-fedramp-dod.json"), crm_doc)
    for baseline, spec in baselines(root).items():
        rules, sources = {}, []
        for pack, slug in spec["packs"]:
            pack_rules = parse_pack(os.path.join(srcdir, "conformance-packs",
                                                 "Operational-Best-Practices-for-" + pack + ".yaml"))
            mapping = parse_mappings(os.path.join(srcdir, "control-mappings", slug + ".html"))
            sources.append({"pack": pack, "pack_url": PACK_URL.format(pack, sha=sha),
                            "control_mapping_url": DOCS.format(slug)})
            for name, r in pack_rules.items():
                cur = rules.setdefault(name, dict(r, controls=[], frameworks=[]))
                cur["frameworks"].append(pack)
                for c in mapping.get(name, []):
                    c = dict(c, framework=pack)
                    if c not in cur["controls"]:
                        cur["controls"].append(c)
        services = {}
        for r in sorted(rules.values(), key=lambda r: r["config_rule"]):
            services.setdefault(_service(r), []).append(r)
        unmapped = sorted(n for n, r in rules.items() if not r["controls"])
        coverage, missing = coverage_for(rules, checked, disp)
        per_baseline[baseline] = coverage
        for n in missing:
            missing_all.setdefault(n, []).append(baseline)
        bcrm = crm_for(baseline, crm, srcdir)
        # Tag each Config-rule control with whether AWS lets the customer inherit it.
        by_key = {}
        for row in bcrm or crm:
            by_key.setdefault(row["control_key"], set()).add(row["can_be_inherited_from_aws"])
        for r in rules.values():
            for c in r["controls"]:
                key = control_key(c["id"])
                inh = by_key.get(key)
                if inh:
                    c["can_be_inherited_from_aws"] = sorted(x for x in inh if x)
                elif bcrm is not None and key:
                    # Proxy packs (full NIST 800-53) map beyond this baseline's control set.
                    c["outside_baseline"] = True
        cat = {
            "baseline": baseline,
            "generated_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "conformance_packs_commit": sha,
            "sources": sources,
            "proxy_note": spec.get("proxy"),
            "rule_count": len(rules),
            "rules_without_control_mapping": unmapped,
            "coverage": coverage,
            "services": dict(sorted(services.items())),
            "artifact_sources": downloaded,
            "crm_note": (CRM_SCOPE.get(baseline) or {}).get("note") or (
                None if bcrm is not None else
                "Overlay baseline: no framework-specific AWS CRM. Responsibility comes from the "
                "FedRAMP base the profile composes it with (see catalog/_crm-fedramp-dod.json)."),
            "crm": bcrm,
            "cmmc_crm": cmmc if baseline == "dod-il4" else None,
            "service_guidance": guidance_for(baseline, guide),
        }
        frag = json.load(open(os.path.join(root, "baselines", baseline + ".json")))
        for svc, entries in inferred_for(baseline, inferred, guide, rules, frag).items():
            cat["service_guidance"].setdefault(svc, {"guide_service": None, "topics": []})
            cat["service_guidance"][svc]["inferred_controls"] = entries
        cat["inferred_controls_note"] = (
            "service_guidance.*.inferred_controls are SGEA-inferred (provenance=inferred), not "
            "AWS-published; attributes validated against provider schema %s. Require ISSO review."
            % schema_providers)
        _write(os.path.join(outdir, baseline + ".json"), _redact(cat))
        if include_artifact_text:
            _write(os.path.join(outdir, "_nda", baseline + ".json"), cat)
        update_vertical(root, baseline, responsibility_section(cat))
        summary.append((baseline, len(rules), len(services), len(unmapped),
                        len(bcrm) if bcrm is not None else 0, bool(spec.get("proxy"))))
    if missing_all:
        raise SystemExit("Config rules with no check and no disposition (add a mapping to "
                         "config-rule-plan-checks.json or an entry to rule-dispositions.yaml):\n  " +
                         "\n  ".join("%s (%s)" % (n, ", ".join(b)) for n, b in sorted(missing_all.items())))
    stale = sorted(set(disp) - {n for cov in per_baseline.values() for n in cov})
    if stale:
        raise SystemExit("rule-dispositions.yaml names rules no composed pack contains: " + ", ".join(stale))
    coverage_report(root, per_baseline, disp)
    return summary


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    ap.add_argument("--include-artifact-text", action="store_true",
                    help="also write full catalogs incl. verbatim AWS Artifact text to catalog/_nda/ (NDA — do not share)")
    args = ap.parse_args(argv)
    print("%-20s %6s %9s %9s %9s  %s" % ("baseline", "rules", "services", "unmapped", "crm rows", "proxy"))
    for b, n, s, u, c, p in build(args.root, args.include_artifact_text):
        print("%-20s %6d %9d %9d %9s  %s" % (b, n, s, u, c or "-", "yes" if p else ""))
    return 0


if __name__ == "__main__":
    sys.exit(main())
